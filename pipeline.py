# -*- coding: utf-8 -*-
"""
Nature Index -> PNU Scholar matching pipeline.

Steps performed here (steps 3-5 of the agent):
  3) For each journal (from FAQ list), collect ISSN / eISSN from independent
     authoritative sources (OpenAlex primary, Crossref secondary).
  4) Search PNU Scholar (scholar.pusan.ac.kr wp-json API) by ISSN (and by name
     as fallback), verify the matched journal is really the same journal
     (ISSN agreement + title match, to guard against same-name journals),
     read whether its Registered DB list contains "Nature Index" (Y/N), and
     record the journal_id + detail URL (https://scholar.pusan.ac.kr/journals/{id}).
  5) Write everything into a single .xlsx.

The FAQ journal list (step 2) is captured separately by agent-browser
(see run.ps1) into a text file and parsed here.

Usage:
  python pipeline.py --faq <faq_text_file> --out <output_dir> [--limit N] [--date YYYYMMDD]
"""
import sys, os, io, re, json, time, argparse, urllib.request, urllib.parse
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Polite-pool contact for OpenAlex/Crossref. Set env NATURE_INDEX_MAILTO to your
# own email for slightly higher rate limits; a neutral default is used otherwise.
MAILTO = os.environ.get("NATURE_INDEX_MAILTO", "nature-index-agent@example.com")
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36"
SCHOLAR_API = "https://scholar.pusan.ac.kr/wp-json/rm/v1/journals"

STOP = {"the", "of", "and", "for", "in", "on", "a", "an", "to"}

# ----------------------------- HTTP with cache -----------------------------
_CACHE = {}
_CACHE_PATH = None

def load_cache(path):
    global _CACHE, _CACHE_PATH
    _CACHE_PATH = path
    if path and os.path.exists(path):
        try:
            _CACHE = json.load(open(path, encoding="utf-8"))
        except Exception:
            _CACHE = {}

def save_cache():
    if _CACHE_PATH:
        tmp = _CACHE_PATH + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(_CACHE, f, ensure_ascii=False)
        os.replace(tmp, _CACHE_PATH)   # atomic: never leaves a half-written cache

def get_json(url, tries=5, pause=0.2):
    if url in _CACHE and _CACHE[url] is not None:
        return _CACHE[url]
    last = None
    for i in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "application/json"})
            with urllib.request.urlopen(req, timeout=40) as r:
                data = json.loads(r.read().decode("utf-8", "replace"))
            _CACHE[url] = data
            time.sleep(pause)
            return data
        except urllib.error.HTTPError as e:
            last = e
            if e.code == 429:            # rate limited -> longer backoff
                time.sleep(2.5 * (i + 1))
            else:
                time.sleep(0.7 * (i + 1))
        except Exception as e:
            last = e
            time.sleep(0.7 * (i + 1))
    print("  ! HTTP fail:", url, "->", last)
    return None   # do NOT cache failures, so re-runs retry them

# ----------------------------- text helpers -----------------------------
def norm(s):
    if not s:
        return ""
    s = s.lower().replace("&", " and ")     # unify '&' and 'and' (scholar mixes them)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()

def strip_the(n):
    return n[4:] if n.startswith("the ") else n

def toks(s):
    return set(t for t in norm(s).split() if t not in STOP)

def overlap(a, b):
    ta, tb = toks(a), toks(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / min(len(ta), len(tb))

def jaccard(a, b):
    ta, tb = toks(a), toks(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)

def valid_issn(x):
    return bool(x) and bool(re.match(r"^\d{4}-\d{3}[\dxX]$", str(x).strip()))

# ----------------------------- FAQ parsing -----------------------------
def parse_faq(path):
    text = open(path, "rb").read().decode("utf-8-sig", "replace")
    anchor = "science journals and conference proceedings:"
    ai = text.find(anchor)
    q13 = text.find("13. What institutions are included", ai if ai > 0 else 0)
    region = text[ai:q13] if (ai > 0 and q13 > ai) else text
    pat = re.compile(r"^(.*\S)\s*\((\d+)\s+articles?\)\s*$")
    out = []
    for line in region.splitlines():
        line = line.strip()
        m = pat.match(line)
        if m:
            out.append({"name": m.group(1).strip(), "articles": int(m.group(2))})
    return out

# ----------------------------- external ISSN -----------------------------
_OA_ENABLED = [True]   # OpenAlex now uses a daily $ budget; disable for the run once exhausted

def openalex_issns(name):
    """Return (issn_set, best_display_name, exact_bool). Optional source (bonus)."""
    if not _OA_ENABLED[0]:
        return set(), "", False
    url = "https://api.openalex.org/sources?search=" + urllib.parse.quote(name) + \
          "&per-page=10&select=id,display_name,issn_l,issn,type&mailto=" + MAILTO
    d = get_json(url, tries=1)
    if d is None:                       # budget/rate error -> stop calling OpenAlex this run
        _OA_ENABLED[0] = False
        print("  (OpenAlex unavailable -> using Crossref + Scholar only)")
        return set(), "", False
    if not d.get("results"):
        return set(), "", False
    nf = strip_the(norm(name))
    exact = None
    for it in d["results"]:
        dn = strip_the(norm(it.get("display_name", "")))
        if dn == nf and (it.get("issn_l") or it.get("issn")):
            exact = it
            break
    pick, is_exact = (exact, True) if exact else (None, False)
    if not pick:
        # best fuzzy among journal/book-series types with an ISSN (display only)
        cand = [it for it in d["results"]
                if it.get("type") in ("journal", "book series", "ebook platform")
                and (it.get("issn_l") or it.get("issn"))]
        if cand:
            pick = max(cand, key=lambda it: jaccard(name, it.get("display_name", "")))
    issns = set()
    if pick:
        if pick.get("issn_l"):
            issns.add(pick["issn_l"])
        for x in (pick.get("issn") or []):
            issns.add(x)
    return {i for i in issns if valid_issn(i)}, (pick.get("display_name") if pick else ""), is_exact

def crossref_issns(name):
    url = "https://api.crossref.org/journals?query=" + urllib.parse.quote(name) + \
          "&rows=20&mailto=" + MAILTO
    d = get_json(url)
    if not d or not d.get("message", {}).get("items"):
        return set(), "", None, None, False
    nf = strip_the(norm(name))
    items = d["message"]["items"]
    exact = None
    for it in items:
        if strip_the(norm(it.get("title", ""))) == nf:
            exact = it
            break
    is_exact = exact is not None
    pick = exact if exact else (items[0] if jaccard(name, items[0].get("title", "")) >= 0.6 else None)
    issns, p_issn, e_issn = set(), None, None
    disp = ""
    if pick:
        disp = pick.get("title", "")
        for x in (pick.get("ISSN") or []):
            if valid_issn(x):
                issns.add(x)
        for t in (pick.get("issn-type") or []):
            if t.get("type") == "print" and valid_issn(t.get("value")):
                p_issn = t["value"]
            elif t.get("type") == "electronic" and valid_issn(t.get("value")):
                e_issn = t["value"]
    return issns, disp, p_issn, e_issn, is_exact

# ----------------------------- scholar -----------------------------
def scholar_search(ks, per=25):
    url = SCHOLAR_API + "?record_per_page=%d&page=1&current_page=1&ks=%s" % (per, urllib.parse.quote(ks))
    d = get_json(url)
    if not d or not d.get("result"):
        return []
    return d["result"].get("records", []) or []

def has_nature_index(rec):
    for db in (rec.get("registered_db") or []):
        if norm(db.get("name", "")) == "nature index":
            return True
    return False

def dbs_str(rec):
    return "/".join(db.get("name", "") for db in (rec.get("registered_db") or []))

def _seq(s):
    """Ordered token list; '&' -> 'and'; keeps stopwords (for phrase containment)."""
    s = (s or "").lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s).strip()
    return s.split()

def _sublist(short, long):
    n = len(short)
    for i in range(len(long) - n + 1):
        if long[i:i + n] == short:
            return True
    return False

def phrase_contained(name, title):
    """True if either title's core (minus leading 'the') is a contiguous phrase of the other.
    Bidirectional so 'The ISME Journal: ...' matches scholar's shorter 'ISME JOURNAL'.
    Guarded to >=2 tokens on the shorter side to avoid single-word over-matching."""
    a, b = _seq(name), _seq(title)
    if a and a[0] == "the":
        a = a[1:]
    if b and b[0] == "the":
        b = b[1:]
    if not a or not b:
        return False
    short, long = (a, b) if len(a) <= len(b) else (b, a)
    if len(short) < 2:
        return False
    return _sublist(short, long)

def title_parts(title):
    """Scholar titles are often 'ACRONYM-FULL NAME' or 'NAME: subtitle'."""
    return [norm(x) for x in re.split(r"[-:]", title or "") if x.strip()]

def score_candidate(name, driver_issns, rec):
    title = rec.get("journal_title", "")
    r_issn, r_eissn = rec.get("issn", ""), rec.get("eissn", "")
    issn_hit = (valid_issn(r_issn) and r_issn in driver_issns) or (valid_issn(r_eissn) and r_eissn in driver_issns)
    nf, nt = strip_the(norm(name)), strip_the(norm(title))
    exact = (nf == nt)
    part_exact = any(strip_the(pp) == nf for pp in title_parts(title))
    contained = phrase_contained(name, title)
    jac = jaccard(name, title)
    extra = len(toks(title) - toks(name))
    score = 0.0
    if issn_hit:
        score += 100
    if exact:
        score += 70
    elif part_exact:
        score += 55
    elif contained:
        score += 35
    score += 15 * jac
    score -= 4 * extra
    return score, issn_hit, exact, part_exact, contained, jac

def _pick(name, driver_issns, cands):
    """Return (rec, method, info) for the best candidate, or (None,'not_found',{})."""
    if not cands:
        return None, "not_found", {}
    scored = []
    for rec in cands.values():
        sig = score_candidate(name, driver_issns, rec)
        scored.append((sig, rec))
    scored.sort(key=lambda x: x[0][0], reverse=True)
    (sc, issn_hit, exact, part_exact, contained, jac), rec = scored[0]
    # acceptance (user rule 1): ANY ISSN/eISSN agreement = confirmed same journal.
    if issn_hit:
        method = "issn_match"
    elif exact:
        method = "title_exact"
    elif part_exact or contained:
        method = "title_match"
    elif jac >= 0.6 and len(toks(name)) >= 2:
        method = "title_fuzzy"
    else:
        return None, "not_found", {"top_title": rec.get("journal_title", ""), "top_score": round(sc, 1)}
    return rec, method, {"issn_hit": issn_hit, "exact": exact, "part": part_exact,
                         "contained": contained, "jaccard": round(jac, 2), "score": round(sc, 1)}

def name_variants(name):
    """Search-query variants to defeat scholar's 'and'/'&' and 'The'/subtitle quirks."""
    out, seen = [], set()
    def add(x):
        x = (x or "").strip()
        if x and x.lower() not in seen:
            seen.add(x.lower()); out.append(x)
    add(name)
    base2 = re.sub(r"^[Tt]he\s+", "", name)
    add(base2)
    for base in (name, base2):
        if ":" in base:
            add(base.split(":")[0])
        low = base.lower()
        if " and " in low:
            add(re.sub(r"\s+and\s+", " & ", base, flags=re.I))
        if "&" in base:
            add(base.replace("&", "and"))
    return out

CONFIDENT = ("issn_match", "title_exact", "title_match")

def match_scholar(name, driver_issns, pool_issns):
    # Lazy: primary name search first; expand with query variants + ISSN lookups if not confident.
    cands = {}
    def add_search(q, per=30):
        for r in scholar_search(q, per=per):
            cands.setdefault(r.get("journal_id"), r)
    add_search(name, 30)
    rec, method, info = _pick(name, driver_issns, cands)
    if method in ("issn_match", "title_exact"):   # only strongest signals short-circuit
        return rec, method, info
    for q in name_variants(name)[1:]:          # variants beyond the primary name
        add_search(q, 30)
    for issn in list(pool_issns)[:5]:
        add_search(issn, 5)
    return _pick(name, driver_issns, cands)

# ----------------------------- main -----------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--faq", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--date", default="")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--cache", default="")
    args = ap.parse_args()

    os.makedirs(args.out, exist_ok=True)
    cache = args.cache or os.path.join(args.out, "_http_cache.json")
    load_cache(cache)

    journals = parse_faq(args.faq)
    print("Parsed FAQ journals:", len(journals))
    if not journals:
        print("ERROR: no journals parsed from FAQ text — aborting.")
        sys.exit(2)
    if args.limit:
        journals = journals[:args.limit]
        print("Limited to:", len(journals))

    # Step 2 deliverable: raw Nature Index journal list from FAQ Q12
    fwb = openpyxl.Workbook(); fws = fwb.active; fws.title = "nature_index_faq"
    fws.append(["no", "journal", "articles(12M)"])
    for k, jj in enumerate(journals, 1):
        fws.append([k, jj["name"], jj["articles"]])
    fws.column_dimensions["B"].width = 60
    faq_xlsx = os.path.join(args.out, "01_nature_index_faq_%s.xlsx" % (args.date or "list"))
    fwb.save(faq_xlsx)
    print("saved FAQ list:", faq_xlsx)

    rows = []
    for i, j in enumerate(journals, 1):
        name = j["name"]
        oa_issns, oa_disp, oa_exact = openalex_issns(name)
        cr_issns, cr_disp, cr_p, cr_e, cr_exact = crossref_issns(name)
        # trusted ISSNs (drive scholar matching) only from EXACT external matches
        driver_issns = set()
        if oa_exact:
            driver_issns |= set(oa_issns)
        if cr_exact:
            driver_issns |= set(cr_issns)
        display_issns = set(oa_issns) | set(cr_issns)
        ext_issns = display_issns
        rec, method, info = match_scholar(name, driver_issns, display_issns)

        if rec:
            r = {
                "no": i, "journal": name,
                "scholar_title": rec.get("journal_title", ""),
                "publisher": rec.get("publisher", ""),
                "issn": rec.get("issn", ""), "eissn": rec.get("eissn", ""),
                "nature_index_db": "Y" if has_nature_index(rec) else "N",
                "journal_id": rec.get("journal_id", ""),
                "detail_url": "https://scholar.pusan.ac.kr/journals/%s" % rec.get("journal_id", ""),
                "all_registered_db": dbs_str(rec),
                "register_yyyymm": rec.get("register_yyyymm", ""),
                "articles_12m": j["articles"],
                "match_method": method,
                "ext_issn": ", ".join(sorted(ext_issns)) if ext_issns else "",
                "issn_verified": "Y" if info.get("issn_hit") else ("N" if driver_issns else "NA"),
            }
        else:
            r = {
                "no": i, "journal": name, "scholar_title": "", "publisher": "",
                "issn": "", "eissn": "",
                "nature_index_db": "not found", "journal_id": "",
                "detail_url": "", "all_registered_db": "",
                "register_yyyymm": "", "articles_12m": j["articles"],
                "match_method": "not_found",
                "ext_issn": ", ".join(sorted(ext_issns)) if ext_issns else "",
                "issn_verified": "NA",
            }
        rows.append(r)
        flag = r["nature_index_db"]
        print("[%3d/%3d] %-52s -> %-40s id=%-7s NI=%-3s (%s) drv=%d%s%s" % (
            i, len(journals), name[:52], r["scholar_title"][:40], r["journal_id"], flag, method,
            len(driver_issns), " oaX" if oa_exact else "", " crX" if cr_exact else ""))
        if i % 10 == 0:
            save_cache()
    save_cache()

    write_xlsx(rows, args.out, args.date)
    # summary
    n_found = sum(1 for r in rows if r["match_method"] != "not_found")
    n_ni = sum(1 for r in rows if r["nature_index_db"] == "Y")
    n_addable = sum(1 for r in rows if r["nature_index_db"] == "N")
    n_nf = sum(1 for r in rows if r["match_method"] == "not_found")
    print("\n=== SUMMARY ===")
    print("total FAQ journals :", len(rows))
    print("matched in scholar :", n_found)
    print("  - tagged NatureIndex (Y):", n_ni)
    print("  - NOT tagged        (N) :", n_addable)
    print("not found in scholar     :", n_nf)

COLS = [
    ("no", "no", 5), ("journal", "Journal (Nature Index FAQ)", 46),
    ("scholar_title", "Scholar Title", 42), ("publisher", "Publisher", 26),
    ("issn", "ISSN", 12), ("eissn", "eISSN", 12),
    ("nature_index_db", "Nature Index DB?", 15),
    ("journal_id", "journal_id", 11),
    ("detail_url", "Scholar URL", 46),
    ("all_registered_db", "Registered DB (all)", 30),
    ("register_yyyymm", "등재년도", 10),
    ("articles_12m", "articles(12M)", 12),
    ("match_method", "match_method", 13),
    ("ext_issn", "ext ISSN (OpenAlex/Crossref)", 26),
    ("issn_verified", "ISSN verified", 12),
]

def _sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(color="FFFFFF", bold=True)
    ws.append([c[1] for c in COLS])
    for ci, c in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=ci)
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = c[2]
    ni_fill = PatternFill("solid", fgColor="C6EFCE")
    add_fill = PatternFill("solid", fgColor="FFEB9C")
    nf_fill = PatternFill("solid", fgColor="F2F2F2")
    for r in rows:
        ws.append([r.get(c[0], "") for c in COLS])
        row_i = ws.max_row
        fill = None
        if r["nature_index_db"] == "Y":
            fill = ni_fill
        elif r["nature_index_db"] == "N":
            fill = add_fill
        elif r["match_method"] == "not_found":
            fill = nf_fill
        if fill:
            ws.cell(row=row_i, column=7).fill = fill
    ws.freeze_panes = "A2"
    return ws

def write_xlsx(rows, out, date):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "result", rows)
    _sheet(wb, "to_add (NI=N)", [r for r in rows if r["nature_index_db"] == "N"])
    _sheet(wb, "not_found", [r for r in rows if r["match_method"] == "not_found"])
    _sheet(wb, "already_NI (NI=Y)", [r for r in rows if r["nature_index_db"] == "Y"])
    fn = "nature_index_scholar_%s.xlsx" % (date or "output")
    path = os.path.join(out, fn)
    wb.save(path)
    print("saved xlsx:", path)

if __name__ == "__main__":
    main()
